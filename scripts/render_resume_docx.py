#!/usr/bin/env python3
"""Create an editable resume DOCX; render and inspect it to confirm pagination."""
import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from render_resume import ORDER, LABELS, validate, prepare_logos, logo_size

BLUE = '2F5496'


def xml(parent, tag, **attributes):
    node = OxmlElement('w:' + tag)
    for key, value in attributes.items():
        node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--font-size', type=float, default=10.5, help='Body size; rounded to a Word half-point')
    parser.add_argument('--line-spacing', type=float, default=14.2)
    parser.add_argument('--paragraph-space', type=float, default=1.2)
    parser.add_argument('--font-family', default='Microsoft YaHei')
    args = parser.parse_args()
    if not 9.8 <= args.font_size <= 11:
        parser.error('Use readable body text between 9.8 and 11 pt')
    args.font_size = round(args.font_size * 2) / 2
    if args.line_spacing < args.font_size + 2:
        parser.error('Line spacing must leave at least 2 pt above the font size')
    if args.paragraph_space < 0:
        parser.error('Paragraph spacing cannot be negative')
    data = json.loads(args.input.read_text(encoding='utf-8'))
    validate(data)
    prepare_logos(data, args.input.parent)
    if data['profile'].get('photo'):
        parser.error('Use the document workflow to place the supplied portrait and avoid text overlap')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Pt(36)
    sec.top_margin, sec.bottom_margin = Pt(32), Pt(32)
    sec.header_distance = sec.footer_distance = Pt(0)
    width = sec.page_width.pt - 72
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(args.font_size)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), args.font_family)
    fmt = normal.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(args.line_spacing)
    fmt.space_before, fmt.space_after = Pt(0), Pt(args.paragraph_space)
    fmt.widow_control = False
    # Prevent the document grid from overriding deliberate line spacing.
    for old in list(sec._sectPr.findall(qn('w:docGrid'))):
        sec._sectPr.remove(old)
    xml(sec._sectPr, 'docGrid', type='default', linePitch='0')
    settings = doc.settings.element
    xml(settings, 'doNotUseHTMLParagraphAutoSpacing')
    # Some installed templates decorate Title with a blue rule. The reference
    # resume has rules only under module headings, never under the name.
    for style in doc.styles:
        for node in list(style.element.findall('.//' + qn('w:pBdr'))):
            node.getparent().remove(node)

    def runs(p, value, bold=False, color=None, size=None):
        for part in re.split(r'(\*\*.*?\*\*)', re.sub(r'\s*\n\s*', ' ', value)):
            if not part:
                continue
            marked = part.startswith('**') and part.endswith('**')
            r = p.add_run(part[2:-2] if marked else part)
            r.bold = bool(bold or marked)
            r.font.name = 'Arial'
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), args.font_family)
            r.font.color.rgb = RGBColor.from_string(color or '000000')
            if size is not None:
                r.font.size = Pt(size)
        return p

    def paragraph(value, label='', keep=False):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = keep
        if label:
            runs(p, label, True)
        runs(p, value)
        return p

    def heading(title):
        p = paragraph('', keep=True)
        runs(p, title, True, BLUE, 12)
        p.paragraph_format.space_before = Pt(2.0)
        p.paragraph_format.space_after = Pt(3.0)
        p.paragraph_format.line_spacing = Pt(15.5)
        xml(xml(p._p.get_or_add_pPr(), 'pBdr'), 'bottom', val='single', sz='12', space='1', color='000000')

    def row(left, center, right, logo=None):
        p = paragraph('', keep=True)
        tabs = p.paragraph_format.tab_stops
        tabs.add_tab_stop(Pt(width / 2), WD_TAB_ALIGNMENT.CENTER)
        tabs.add_tab_stop(Pt(width), WD_TAB_ALIGNMENT.RIGHT)
        runs(p, left + '\t', True)
        if logo:
            image_width, image_height = logo_size(logo)
            picture = p.add_run().add_picture(logo, width=Pt(image_width), height=Pt(image_height))
            picture._inline.docPr.set('descr', center + ' company logo')
            runs(p, ' ', True)
        runs(p, center + '\t' + right, True)
        return p

    profile = data['profile']
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = Pt(25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    runs(p, profile['name'], True, BLUE, 20)
    for line in profile.get('contact_lines', []):
        p = paragraph('', keep=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = Pt(14)
        p.paragraph_format.space_after = Pt(0.7)
        for part in re.split(r'(\*\*.*?\*\*)', line):
            marked = part.startswith('**') and part.endswith('**')
            runs(p, part[2:-2] if marked else part, marked, BLUE if marked else None, 10.5)
    if data.get('education'):
        heading('教育背景')
        for school in data['education']:
            row(school['institution'], school['degree'], school.get('dates', ''))
            for detail in school.get('details', []):
                paragraph(detail)
        if data.get('honors'):
            paragraph(data['honors'], '荣誉奖项：')
    sections = {s['title']: s for s in data.get('sections', [])}
    for title in ORDER:
        if title not in sections:
            continue
        heading(title)
        for entry in sections[title]['entries']:
            row(entry.get('dates', ''), entry['name'], entry.get('role', ''), entry.get('company_logo'))
            for key, label in LABELS:
                paragraph(entry[key], label, keep=key != 'results')
    summary = data.get('summary', {})
    if any(summary.get(k) for k in ('skills', 'achievements', 'evaluation')):
        heading('个人技能与自我评价')
        for key, label in [('skills', '个人技能：'), ('achievements', '成果总结：'), ('evaluation', '自我评价：')]:
            if summary.get(key):
                paragraph(summary[key], label)
    doc.core_properties.title = profile['name'] + ' 个人简历'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(json.dumps({'output': str(args.output), 'font_size': args.font_size,
                      'editable': True, 'pagination': 'Confirm by rendering all pages'}, ensure_ascii=False))


if __name__ == '__main__':
    main()
